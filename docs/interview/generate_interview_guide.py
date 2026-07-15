import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/interview/高级前端开发工程师面试指导.md"
TARGET = ROOT / "docs/interview/高级前端开发工程师面试指导.docx"

CJK_FONT = "Arial Unicode MS"
LATIN_FONT = "Arial Unicode MS"
BODY_COLOR = RGBColor(36, 43, 51)
MUTED_COLOR = RGBColor(91, 103, 114)
ACCENT_COLOR = RGBColor(14, 92, 102)
ACCENT_DARK = RGBColor(10, 66, 74)
QUESTION_FILL = "E8F2F3"
CALLOUT_FILL = "F3F6F7"
RULE_COLOR = "80AEB4"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, size: float, bold: bool = False, color=BODY_COLOR, italic: bool = False) -> None:
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 0, line: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_keep(paragraph, keep_next: bool = False, keep_lines: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines
    paragraph.paragraph_format.widow_control = True


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, side: str, color: str, size: str = "8", space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)
    borders.append(border)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, 8.5, color=MUTED_COLOR)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY_COLOR
    fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(4)

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(10.2)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.left_indent = Cm(0.65)
        style.paragraph_format.first_line_indent = Cm(-0.28)
        style.paragraph_format.line_spacing = 1.18
        style.paragraph_format.space_after = Pt(2.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(hp, 0, 0)
    set_run_font(hp.add_run("高级前端开发工程师面试指导  |  2026-07"), 8.2, color=MUTED_COLOR)
    set_paragraph_border(hp, "bottom", RULE_COLOR, size="4", space="2")

    footer = section.footer
    fp = footer.paragraphs[0]
    set_paragraph_spacing(fp, 0, 0)
    add_page_number(fp)

    document.core_properties.title = "高级前端开发工程师面试指导"
    document.core_properties.subject = "Senior Frontend Engineer / Web3 Frontend Engineer 面试复习"
    document.core_properties.author = "Quest Price"


def add_inline_markdown(paragraph, text: str, size: float = 10.5, color=BODY_COLOR, bold: bool = False) -> None:
    text = text.rstrip().replace("  ", " ")
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        is_url = part.startswith("http://") or part.startswith("https://")
        clean = part[2:-2] if is_bold else part[1:-1] if is_code else part
        run = paragraph.add_run(clean)
        set_run_font(
            run,
            size - 0.2 if is_url else size,
            bold=bold or is_bold or is_code,
            color=ACCENT_DARK if is_url else color,
        )
        if is_url:
            run.font.underline = True


def add_cover(document: Document, lines: list[str]) -> int:
    title = lines[0][2:].strip()
    metadata = []
    index = 1
    while index < len(lines) and not lines[index].startswith("## "):
        if lines[index].strip():
            metadata.append(lines[index].strip())
        index += 1

    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, before=72, after=0)

    accent = document.add_paragraph()
    accent.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(accent, 0, 16)
    set_run_font(accent.add_run("SENIOR FRONTEND · INTERVIEW PLAYBOOK"), 10.5, bold=True, color=ACCENT_COLOR)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 0, 12)
    set_run_font(p.add_run(title), 27, bold=True, color=ACCENT_DARK)

    rule = document.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(rule, 0, 22)
    set_paragraph_border(rule, "bottom", RULE_COLOR, size="16", space="1")

    for item in metadata:
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(meta, 0, 5)
        add_inline_markdown(meta, item, 11, color=MUTED_COLOR)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note, before=36, after=0, line=1.25)
    set_paragraph_shading(note, CALLOUT_FILL)
    note.paragraph_format.left_indent = Cm(1.4)
    note.paragraph_format.right_indent = Cm(1.4)
    add_inline_markdown(note, "基于最终简历、已确认项目事实与当前源码证据整理。用于口述复习，不作为逐字背诵稿。", 10.2, color=BODY_COLOR)

    document.add_page_break()
    return index


def add_contents(document: Document, headings: list[str]) -> None:
    p = document.add_paragraph()
    set_paragraph_spacing(p, 0, 8)
    set_keep(p, keep_next=True)
    set_run_font(p.add_run("目录"), 18, bold=True, color=ACCENT_DARK)
    set_paragraph_border(p, "bottom", RULE_COLOR, size="8", space="4")

    rows = (len(headings) + 1) // 2
    table = document.add_table(rows=rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Cm(8.5)
        row.cells[1].width = Cm(8.5)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F8FAFA")
            cell.margin_top = Cm(0.08)
            cell.margin_bottom = Cm(0.08)
    for idx, heading in enumerate(headings):
        row = idx % rows
        col = idx // rows
        cell = table.cell(row, col)
        cp = cell.paragraphs[0]
        set_paragraph_spacing(cp, 2.5, 2.5)
        set_run_font(cp.add_run(heading), 9.5, bold=True, color=ACCENT_DARK)

    hint = document.add_paragraph()
    set_paragraph_spacing(hint, 8, 8, 1.15)
    set_paragraph_shading(hint, CALLOUT_FILL)
    hint.paragraph_format.left_indent = Cm(0.3)
    hint.paragraph_format.right_indent = Cm(0.3)
    add_inline_markdown(hint, "导航提示：DOCX 保留 Heading 层级，可在 Word 或 LibreOffice 的导航窗格中按章节和题号跳转。", 9.5, color=MUTED_COLOR)


def add_section_heading(document: Document, text: str, first_section: bool) -> None:
    if not first_section:
        document.add_page_break()
    p = document.add_paragraph()
    set_keep(p, keep_next=True)
    set_paragraph_spacing(p, 0, 9)
    set_run_font(p.add_run(text), 17, bold=True, color=ACCENT_DARK)
    set_paragraph_border(p, "bottom", RULE_COLOR, size="10", space="4")
    p.style = document.styles["Heading 1"]
    for run in p.runs:
        set_run_font(run, 17, bold=True, color=ACCENT_DARK)


def add_question_heading(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.style = document.styles["Heading 2"]
    set_keep(p, keep_next=True)
    set_paragraph_spacing(p, 9, 5, 1.08)
    set_paragraph_shading(p, QUESTION_FILL)
    p.paragraph_format.left_indent = Cm(0.22)
    p.paragraph_format.right_indent = Cm(0.18)
    set_paragraph_border(p, "left", RULE_COLOR, size="18", space="6")
    add_inline_markdown(p, text, 11.2, color=ACCENT_DARK, bold=True)


def add_label(document: Document, text: str) -> None:
    label_colors = {
        "面试官问": ACCENT_DARK,
        "参考回答": ACCENT_COLOR,
        "可能追问": RGBColor(105, 75, 20),
        "回答要点": RGBColor(63, 91, 54),
        "容易踩坑": RGBColor(145, 64, 55),
    }
    p = document.add_paragraph()
    set_keep(p, keep_next=True)
    set_paragraph_spacing(p, 5, 2)
    set_run_font(p.add_run(text), 9.5, bold=True, color=label_colors.get(text, ACCENT_DARK))


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    set_keep(p, keep_lines=False)
    set_paragraph_spacing(p, 0, 4, 1.22)
    add_inline_markdown(p, text, 10.5)


def add_bullet(document: Document, text: str, numbered: bool = False) -> None:
    p = document.add_paragraph(style="List Number" if numbered else "List Bullet")
    set_keep(p, keep_lines=False)
    set_paragraph_spacing(p, 0, 2.5, 1.18)
    add_inline_markdown(p, text, 10.2)


def build_document(source: Path, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    start = add_cover(document, lines)
    section_headings = [line[3:].strip() for line in lines if line.startswith("## ")]
    add_contents(document, section_headings)

    first_section = True
    for raw_line in lines[start:]:
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            add_section_heading(document, line[3:], first_section)
            first_section = False
        elif line.startswith("### Q"):
            add_question_heading(document, line[4:])
        elif line.startswith("### "):
            add_question_heading(document, line[4:])
        elif re.fullmatch(r"\*\*[^*]+\*\*", line):
            add_label(document, line[2:-2])
        elif line.startswith("- "):
            add_bullet(document, line[2:])
        elif re.match(r"^\d+\. ", line):
            add_bullet(document, re.sub(r"^\d+\. ", "", line), numbered=True)
        else:
            add_body(document, line)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


if __name__ == "__main__":
    build_document(SOURCE, TARGET)
