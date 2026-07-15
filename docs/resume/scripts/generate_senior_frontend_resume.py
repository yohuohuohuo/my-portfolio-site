import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "docs/resume/output/高级前端开发工程师-中文主简历-202607.md"
TARGET = ROOT / "docs/resume/output/高级前端开发工程师-中文主简历-202607.docx"
CJK_FONT = "Hiragino Sans GB"
PAGE_BREAK_SECTION = "精选项目"
LATIN_FONT = "Arial"
BODY_COLOR = RGBColor(32, 39, 48)
ACCENT_COLOR = RGBColor(15, 78, 92)


def set_run_font(run, size: float, bold: bool = False, color=BODY_COLOR) -> None:
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)


def set_paragraph_spacing(
    paragraph, before: float, after: float, line: float = 1.0
) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.50)
    section.bottom_margin = Cm(1.50)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    bullet = document.styles["List Bullet"]
    bullet.font.name = LATIN_FONT
    bullet.font.size = Pt(11)
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    bullet.paragraph_format.left_indent = Cm(0.45)
    bullet.paragraph_format.first_line_indent = Cm(-0.20)


def add_inline_markdown(paragraph, text: str, size: float, color=BODY_COLOR) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2), size, color)
            continue
        bold = part.startswith("**") and part.endswith("**")
        inline_code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else part[1:-1] if inline_code else part
        run = paragraph.add_run(clean)
        set_run_font(run, size, bold=bold or inline_code, color=color)


def add_hyperlink(paragraph, text: str, url: str, size: float, color) -> None:
    relationship_id = paragraph.part.relate_to(
        url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)

    run = Run(run_element, paragraph)
    run.text = text
    set_run_font(run, size, color=color)
    run.font.underline = True


def add_name(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, 0, 1.5)
    set_run_font(paragraph.add_run(text), 22, bold=True, color=ACCENT_COLOR)


def add_role(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, 0, 3.5)
    set_run_font(paragraph.add_run(text), 12, bold=True, color=BODY_COLOR)


def add_contact(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, 0, 0.75)
    add_inline_markdown(paragraph, text, 9.5)


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    if PAGE_BREAK_SECTION and text == PAGE_BREAK_SECTION:
        paragraph.paragraph_format.page_break_before = True
    set_paragraph_spacing(paragraph, 10, 3)
    set_run_font(paragraph.add_run(text), 13, bold=True, color=ACCENT_COLOR)


def add_entry_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_spacing(paragraph, 5, 2)
    add_inline_markdown(paragraph, text, 11.5, color=BODY_COLOR)
    for run in paragraph.runs:
        run.bold = True


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.widow_control = True
    set_paragraph_spacing(paragraph, 0, 2.3, 1.15)
    add_inline_markdown(paragraph, text, 11)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, 0, 2.3, 1.15)
    add_inline_markdown(paragraph, text, 11)


def build_document(
    source: Path,
    target: Path,
    title: str = "高级前端开发工程师中文主简历",
    subject: str = "Senior Frontend Engineer / Web3 Frontend Engineer",
    cjk_font: str = CJK_FONT,
    page_break_section: str = PAGE_BREAK_SECTION,
) -> None:
    global CJK_FONT, PAGE_BREAK_SECTION
    previous_cjk_font = CJK_FONT
    previous_page_break_section = PAGE_BREAK_SECTION
    CJK_FONT = cjk_font
    PAGE_BREAK_SECTION = page_break_section
    document = Document()
    configure_document(document)
    document.core_properties.title = title
    document.core_properties.subject = subject

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            add_name(document, line[2:])
        elif line.startswith("## "):
            add_section_heading(document, line[3:])
        elif line.startswith("### "):
            add_entry_heading(document, line[4:])
        elif line.startswith("- "):
            add_bullet(document, line[2:])
        elif line.startswith("**") and line.endswith("**"):
            add_role(document, line[2:-2])
        elif line.startswith(("成都 |", "GitHub /", "[GitHub /")):
            add_contact(document, line)
        else:
            add_body(document, line)

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
    finally:
        CJK_FONT = previous_cjk_font
        PAGE_BREAK_SECTION = previous_page_break_section


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate an ATS-friendly resume DOCX")
    parser.add_argument("source", nargs="?", type=Path, default=SOURCE)
    parser.add_argument("target", nargs="?", type=Path, default=TARGET)
    parser.add_argument("--title", default="高级前端开发工程师中文主简历")
    parser.add_argument(
        "--subject", default="Senior Frontend Engineer / Web3 Frontend Engineer"
    )
    parser.add_argument("--cjk-font", default=CJK_FONT)
    parser.add_argument("--page-break-section", default=PAGE_BREAK_SECTION)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_document(
        args.source,
        args.target,
        title=args.title,
        subject=args.subject,
        cjk_font=args.cjk_font,
        page_break_section=args.page_break_section,
    )
